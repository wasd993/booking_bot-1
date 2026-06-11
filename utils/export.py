"""
Экспорт записей в .xlsx через openpyxl.
Возвращает bytes — можно сразу передать в BufferedInputFile.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="2F5496")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_EVEN_FILL   = PatternFill("solid", fgColor="DCE6F1")
_CENTER      = Alignment(horizontal="center", vertical="center")
_LEFT        = Alignment(horizontal="left",   vertical="center")

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

COLUMNS = [
    ("ID",              8,  "id"),
    ("Имя",            22,  "full_name"),
    ("Username",       18,  "username"),
    ("Дата",           14,  "date"),
    ("Время",          10,  "time"),
    ("⭐️ Звёзды",      13,  "stars_paid"),
    ("Статус",         13,  "status"),
    ("Создано",        20,  "created_at"),
]

STATUS_RU = {
    "active":    "активна",
    "cancelled": "отменена",
    "pending":   "ожидает",
}


def _fmt_date(iso: str) -> str:
    try:
        y, m, d = iso.split("-")
        return f"{d}.{m}.{y}"
    except Exception:
        return iso


async def export_appointments_to_excel(appointments: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Записи"

    # ── заголовки ──────────────────────────────
    ws.row_dimensions[1].height = 22
    for col_idx, (header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = _HEADER_FILL
        cell.font      = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border    = _BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── данные ─────────────────────────────────
    for row_idx, apt in enumerate(appointments, start=2):
        even = (row_idx % 2 == 0)
        fill = _EVEN_FILL if even else None

        values = [
            apt.get("id", ""),
            apt.get("full_name") or "—",
            f"@{apt['username']}" if apt.get("username") else "—",
            _fmt_date(apt.get("date", "")),
            apt.get("time", ""),
            apt.get("stars_paid", 0),
            STATUS_RU.get(apt.get("status", ""), apt.get("status", "")),
            str(apt.get("created_at", ""))[:19],
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border    = _BORDER
            cell.alignment = _CENTER if col_idx in (1, 5, 6) else _LEFT
            if fill:
                cell.fill = fill

        ws.row_dimensions[row_idx].height = 18

    # ── итого ──────────────────────────────────
    last = len(appointments) + 2
    ws.cell(row=last, column=5, value="ИТОГО:").font      = Font(bold=True)
    ws.cell(row=last, column=5).alignment                  = _CENTER
    ws.cell(row=last, column=6, value=sum(a.get("stars_paid", 0) for a in appointments))
    ws.cell(row=last, column=6).font                       = Font(bold=True)
    ws.cell(row=last, column=6).alignment                  = _CENTER

    # ── freeze header ──────────────────────────
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
